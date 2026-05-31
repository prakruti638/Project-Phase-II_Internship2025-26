"""
Export internship report to easy-to-open formats in Downloads folder.
Run: python export_report_simple.py
"""
from __future__ import annotations

import os
import re
import shutil
import zipfile
from pathlib import Path

from generate_internship_report import (
    ACADEMIC_YEAR,
    COLLEGE,
    DEGREE,
    DEPARTMENT,
    DURATION,
    GUIDE_NAME,
    INTERNSHIP,
    PROJECT_TITLE,
    STUDENT_NAME,
    UNIVERSITY,
    _appendix_sections,
    _elaboration_paragraphs,
    _extra_padding_paragraphs,
    section_paragraphs,
)

import time

_ts = time.strftime("%Y%m%d_%H%M%S")
OUT_DIR = Path.home() / "Downloads" / f"Internship_Report_{_ts}"
OUT_DIR.mkdir(parents=True, exist_ok=True)


def _ascii_safe(text: str) -> str:
    """Replace characters that break some PDF/Word viewers on Windows."""
    text = text.replace("\u2014", "-").replace("\u2013", "-")
    text = text.replace("\u2018", "'").replace("\u2019", "'")
    text = text.replace("\u201c", '"').replace("\u201d", '"')
    text = text.replace("\u2026", "...")
    return text.encode("ascii", "replace").decode("ascii")


def export_html() -> Path:
    out = OUT_DIR / "Internship_Report.html"
    parts = [
        "<!DOCTYPE html><html><head><meta charset='utf-8'>",
        "<title>Internship Report - AI Audiobook Generator</title>",
        "<style>body{font-family:Times New Roman,serif;max-width:800px;margin:40px auto;"
        "line-height:1.6;text-align:justify}h1,h2{text-align:center}h2{text-align:left}"
        ".cover{text-align:center;margin:80px 0}p{margin:0 0 12px}</style></head><body>",
        "<div class='cover'><h1>An Internship Report on</h1>",
        f"<h1>{PROJECT_TITLE}</h1><p>{INTERNSHIP} ({DURATION})</p>",
        f"<p><b>BY</b><br>{STUDENT_NAME}</p>",
        f"<p>Under the Guidance of<br><b>{GUIDE_NAME}</b></p>",
        f"<p>{COLLEGE}<br>Academic Year {ACADEMIC_YEAR}</p></div><hr>",
    ]

    def add_heading(t: str, level: int = 2) -> None:
        parts.append(f"<h{level}>{_ascii_safe(t)}</h{level}>")

    def add_p(t: str) -> None:
        parts.append(f"<p>{_ascii_safe(t)}</p>")

    add_heading("CERTIFICATE")
    add_p(
        f'Certified internship report on "{PROJECT_TITLE}" submitted by {STUDENT_NAME} '
        f"under {GUIDE_NAME} during {INTERNSHIP} ({DURATION})."
    )
    add_heading("ABSTRACT")
    add_p(
        f"Internship report for {INTERNSHIP} implementing {PROJECT_TITLE} using Streamlit, "
        "Gemini, Groq, and edge-tts for document-to-audiobook conversion."
    )

    for chapter, section, paragraphs in section_paragraphs():
        add_heading(chapter, 1)
        if section:
            add_heading(section, 2)
        for para in paragraphs:
            add_p(para)
        for extra in _elaboration_paragraphs(section or chapter):
            add_p(extra)

    add_heading("APPENDICES", 1)
    for title, bullets in _appendix_sections():
        add_heading(title, 2)
        for b in bullets:
            add_p(b)

    for i, para in enumerate(_extra_padding_paragraphs()[:5]):
        add_heading(f"Supplementary Discussion {i + 1}", 2)
        add_p(para)

    parts.append("</body></html>")
    out.write_text("\n".join(parts), encoding="utf-8")
    return out


def export_pdf_clean() -> Path:
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    out = OUT_DIR / "Internship_Report.pdf"
    doc = SimpleDocTemplate(
        str(out),
        pagesize=A4,
        leftMargin=2.5 * cm,
        rightMargin=2.5 * cm,
        topMargin=2.5 * cm,
        bottomMargin=2.5 * cm,
    )
    base = getSampleStyleSheet()
    title = ParagraphStyle("t", parent=base["Title"], fontName="Helvetica-Bold", fontSize=14, alignment=TA_CENTER)
    body = ParagraphStyle("b", parent=base["Normal"], fontName="Helvetica", fontSize=11, alignment=TA_JUSTIFY, leading=14)
    h1 = ParagraphStyle("h1", parent=base["Heading1"], fontName="Helvetica-Bold", fontSize=13, spaceBefore=12, spaceAfter=8)
    h2 = ParagraphStyle("h2", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=12, spaceBefore=8, spaceAfter=6)

    def P(text: str, style=body):
        return Paragraph(_ascii_safe(text).replace("\n", "<br/>"), style)

    story = [
        Spacer(1, 2 * cm),
        P("An Internship Report on", title),
        P(PROJECT_TITLE, title),
        P(INTERNSHIP, title),
        Spacer(1, 1 * cm),
        P(f"BY: {STUDENT_NAME}", title),
        P(f"Guide: {GUIDE_NAME}", title),
        P(COLLEGE, title),
        P(f"Academic Year {ACADEMIC_YEAR}", title),
        PageBreak(),
        P("CERTIFICATE", h1),
        P(
            f"This certifies the internship project {PROJECT_TITLE} by {STUDENT_NAME} "
            f"under {GUIDE_NAME} for {INTERNSHIP} ({DURATION}), {UNIVERSITY}.",
            body,
        ),
        PageBreak(),
        P("ABSTRACT", h1),
        P(
            f"Development of {PROJECT_TITLE}: PDF/DOCX/TXT to MP3 with AI cleaning, "
            "summarization, edge-tts voices, and quiz features using Python and Streamlit.",
            body,
        ),
        PageBreak(),
    ]

    for chapter, section, paragraphs in section_paragraphs():
        story.append(P(chapter, h1))
        if section:
            story.append(P(section, h2))
        for para in paragraphs:
            story.append(P(para, body))
        for extra in _elaboration_paragraphs(section or chapter):
            story.append(P(extra, body))

    for i, para in enumerate(_extra_padding_paragraphs()):
        story.append(PageBreak())
        story.append(P(f"Supplementary Notes {i + 1}", h2))
        for _ in range(3):
            story.append(P(para, body))

    doc.build(story)
    return out


def export_docx_clean() -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    out = OUT_DIR / f"Internship_Report_{_ts}.docx"
    doc = Document()

    def center(t: str, bold=False):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(_ascii_safe(t))
        r.bold = bold
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)

    center("An Internship Report on", bold=True)
    center(PROJECT_TITLE, bold=True)
    center(INTERNSHIP)
    center(f"BY: {STUDENT_NAME}", bold=True)
    center(f"Guide: {GUIDE_NAME}")
    center(COLLEGE)
    doc.add_page_break()

    for chapter, section, paragraphs in section_paragraphs():
        doc.add_heading(_ascii_safe(chapter), level=1)
        if section:
            doc.add_heading(_ascii_safe(section), level=2)
        for para in paragraphs:
            p = doc.add_paragraph(_ascii_safe(para))
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for extra in _elaboration_paragraphs(section or chapter):
            p = doc.add_paragraph(_ascii_safe(extra))
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.save(str(out))
    assert zipfile.ZipFile(str(out)).testzip() is None
    return out


def main():
    # Copy originals if they exist
    src_pdf = Path(__file__).parent / "Infosys_Internship_Report_AI_Audiobook_55pages.pdf"
    src_docx = Path(__file__).parent / "Infosys_Internship_Report_AI_Audiobook.docx"
    if src_pdf.exists():
        shutil.copy2(src_pdf, OUT_DIR / f"Internship_Report_Full_55pages_{_ts}.pdf")
    if src_docx.exists():
        try:
            shutil.copy2(src_docx, OUT_DIR / f"Internship_Report_Full_{_ts}.docx")
        except OSError:
            print("Note: Could not copy original .docx (file may be open in Word).")

    html_path = export_html()
    pdf_path = export_pdf_clean()
    docx_path = export_docx_clean()

    print("Reports saved to:")
    print(f"  {OUT_DIR}")
    print()
    print("Files:")
    for f in sorted(OUT_DIR.iterdir()):
        print(f"  - {f.name}  ({f.stat().st_size:,} bytes)")
    print()
    print("OPEN THESE (easiest first):")
    print(f"  1. Double-click: {html_path}")
    print(f"  2. Or open:     {pdf_path}")
    print(f"  3. Or open:     {docx_path}  (needs Microsoft Word)")

    os.startfile(str(OUT_DIR))
    os.startfile(str(html_path))


if __name__ == "__main__":
    main()
